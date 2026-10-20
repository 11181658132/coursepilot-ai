"""Extract page/slide-local text. A chunk never crosses a source location."""

import hashlib
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import List, Optional
from urllib.parse import unquote_plus

from app.services.library import Document

PARSER_VERSION = "page-chunks-v1"


@dataclass(frozen=True)
class TextUnit:
    page: int
    locator_type: str
    text: str


@dataclass(frozen=True)
class Chunk:
    chunk_id: str
    path: str
    filename: str
    course: str
    material_type: str
    page: int
    locator_type: str
    week: Optional[int]
    ordinal: int
    text: str

    def as_dict(self) -> dict:
        return asdict(self)


def extract_units(path: Path) -> List[TextUnit]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("Encrypted PDF requires a password")
        return [TextUnit(i, "page", page.extract_text() or "")
                for i, page in enumerate(reader.pages, 1)]
    if suffix == ".pptx":
        from pptx import Presentation

        units = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            parts = []

            def read_shapes(shapes, output):
                for shape in shapes:
                    if hasattr(shape, "shapes"):
                        read_shapes(shape.shapes, output)
                    if shape.has_text_frame:
                        output.append(shape.text_frame.text)
                    if shape.has_table:
                        output.extend(" | ".join(cell.text for cell in row.cells)
                                      for row in shape.table.rows)

            read_shapes(slide.shapes, parts)
            units.append(TextUnit(i, "slide", "\n".join(parts)))
        return units
    if suffix == ".docx":
        from docx import Document as WordDocument

        document = WordDocument(path)
        # DOCX pagination depends on the renderer: report sections, never fake pages.
        text = "\n".join([p.text for p in document.paragraphs] + [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables for row in table.rows
        ])
        return [TextUnit(1, "section", text)]
    if suffix in {".txt", ".md"}:
        return [TextUnit(1, "section", path.read_text(encoding="utf-8-sig"))]
    raise ValueError(f"Unsupported parser: {suffix}")


def chunk_document(document: Document, size: int = 1600, overlap: int = 200) -> List[Chunk]:
    if size < 100 or overlap < 0 or overlap >= size:
        raise ValueError("Chunk size must be >= 100 and 0 <= overlap < size")
    week_match = re.search(r"(?:week|wk)[\s_.-]*(\d{1,2})(?!\d)",
                           unquote_plus(document.relative_path), re.IGNORECASE)
    week = int(week_match[1]) if week_match else None
    chunks = []
    for unit in extract_units(document.absolute_path):
        text = unit.text.replace("\x00", "").strip()
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                # Prefer a paragraph/sentence/word boundary in the last quarter.
                boundaries = [text.rfind(mark, start + size * 3 // 4, end)
                              for mark in ("\n", ". ", "。", " ")]
                boundary = max(boundaries)
                if boundary > start:
                    end = boundary + 1
            content = text[start:end].strip()
            if content:
                identity = f"{document.relative_path}:{document.sha256}:{unit.page}:{start}"
                identity += f":{PARSER_VERSION}:{size}:{overlap}"
                chunks.append(Chunk(
                    hashlib.sha256(identity.encode()).hexdigest()[:32],
                    document.relative_path, document.absolute_path.name, document.course,
                    document.material_type, unit.page, unit.locator_type, week,
                    len(chunks), content,
                ))
            if end == len(text):
                break
            start = max(start + 1, end - overlap)
    if not chunks:
        raise ValueError("No extractable text; scanned/image-only files need OCR")
    return chunks
