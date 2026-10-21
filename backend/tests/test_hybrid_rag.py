import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from app.config import Settings, get_settings
from app.main import app, get_hybrid_base, get_knowledge_base
from app.services.hybrid_kb import HybridKnowledgeBase
from app.services.library import discover_documents
from app.services.parsing import TextUnit, chunk_document, extract_units
from app.services.retrieval import bm25, cosine, rank_candidates
from fastapi.testclient import TestClient


class FakeClient:
    def __init__(self):
        self.embeddings = SimpleNamespace(create=self.embed)
        self.responses = SimpleNamespace(create=self.respond)
        self.embed_calls = 0
        self.model_calls = 0
        self.invalid_quote = False
        self.invalid_id = False
        self.invalid_rerank = False
        self.low_relevance = False
        self.embedding_failure = False

    def embed(self, **kwargs):
        self.embed_calls += 1
        if self.embedding_failure:
            raise RuntimeError("Simulated provider failure")
        vectors = []
        for i, text in enumerate(kwargs["input"]):
            vector = [1.0, 0.0] if "database" in text.lower() else [0.0, 1.0]
            vectors.append(SimpleNamespace(index=i, embedding=vector))
        return SimpleNamespace(data=vectors, usage=SimpleNamespace(total_tokens=100))

    def respond(self, **kwargs):
        self.model_calls += 1
        payload = json.loads(kwargs["input"])
        if kwargs["text"]["format"]["name"] == "relevance_scores":
            scores = [{"chunk_id": item["chunk_id"], "score": 0 if self.low_relevance else 3}
                      for item in payload["passages"]]
            if self.invalid_rerank:
                scores[0]["chunk_id"] = "invented"
            result = {"scores": scores}
        else:
            first = payload["passages"][0]
            source_id = "S99" if self.invalid_id else first["source_id"]
            quote = "this quote was invented" if self.invalid_quote else first["text"][:80]
            result = {"answer": f"Database transactions preserve consistency. [{source_id}]",
                      "citations": [{"source_id": source_id, "quote": quote}]}
        return SimpleNamespace(output_text=json.dumps(result),
                               usage=SimpleNamespace(input_tokens=100, output_tokens=20))


class HybridRagTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.config = Settings(_env_file=None, knowledge_base_dir=root / "knowledge",
                               rag_db_path=root / "rag.sqlite3", rag_engine="hybrid",
                               conversation_db_path=root / "chat.db",
                               index_state_path=root / "state.json",
                               telemetry_path=root / "events.jsonl")
        self.path = self.write("A/lecture/Week 3.md", "Database transactions preserve consistency.")
        self.client = FakeClient()
        self.base = HybridKnowledgeBase(self.config, self.client)

    def tearDown(self):
        app.dependency_overrides.clear()
        self.temp.cleanup()

    def write(self, path, text):
        target = self.config.knowledge_base_dir / path
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding="utf-8")
        return target

    def test_page_local_chunking_and_stable_identity(self):
        document = discover_documents(self.config.knowledge_base_dir)[0]
        units = [TextUnit(1, "page", "First page. " * 40),
                 TextUnit(2, "page", "Second page. " * 40)]
        with patch("app.services.parsing.extract_units", return_value=units):
            chunks = chunk_document(document, 160, 30)
            repeated = chunk_document(document, 160, 30)
        self.assertEqual(chunks, repeated)
        self.assertTrue(all(len(chunk.text) <= 160 for chunk in chunks))
        self.assertTrue(all(chunk.week == 3 for chunk in chunks))
        self.assertTrue(all("Second" not in c.text for c in chunks if c.page == 1))
        self.assertEqual({c.page for c in chunks}, {1, 2})

    def test_week_metadata_decodes_canvas_download_names(self):
        self.write("A/lecture/Course+-+Week+5+-+Notes.md", "Database transactions")
        document = next(d for d in discover_documents(self.config.knowledge_base_dir)
                        if "Week+5" in d.relative_path)
        self.assertEqual(chunk_document(document)[0].week, 5)

    def test_incremental_replace_and_prune_deleted_files(self):
        report = self.base.sync(self.config.knowledge_base_dir)
        self.assertEqual(report.uploaded, 1)
        self.assertGreater(report.embedding_tokens, 0)
        calls = self.client.embed_calls
        self.assertEqual(self.base.sync(self.config.knowledge_base_dir).skipped, 1)
        self.assertEqual(calls, self.client.embed_calls)
        old_id = self.base.retrieve("A", "lecture", "database", rerank=False)[0]["chunk_id"]
        self.path.write_text("Database transactions now include rollback.")
        self.assertEqual(self.base.sync(self.config.knowledge_base_dir).replaced, 1)
        self.assertIsNone(self.base.source(old_id))
        self.path.unlink()
        self.assertEqual(self.base.sync(self.config.knowledge_base_dir).removed, 1)
        self.assertEqual(self.base.manifest(), {})

    def test_preview_is_free_and_does_not_change_index(self):
        report = self.base.sync(self.config.knowledge_base_dir, dry_run=True)
        self.assertEqual(report.uploaded, 1)
        self.assertEqual(self.client.embed_calls, 0)
        self.assertEqual(self.base.manifest(), {})

    def test_embedding_budget_rejects_document_before_provider_call(self):
        config = self.config.model_copy(update={"rag_index_max_embedding_tokens": 1})
        base = HybridKnowledgeBase(config, self.client)
        report = base.sync(config.knowledge_base_dir)
        self.assertEqual(report.failed, 1)
        self.assertEqual(self.client.embed_calls, 0)
        self.assertEqual(base.manifest(), {})

    def test_failed_replacement_preserves_previous_index_but_never_serves_stale_text(self):
        self.base.sync(self.config.knowledge_base_dir)
        before = self.base.manifest()
        self.path.write_text("Changed database content")
        self.client.embedding_failure = True
        report = self.base.sync(self.config.knowledge_base_dir)
        self.assertEqual(report.failed, 1)
        self.assertEqual(before, self.base.manifest())
        with self.assertRaises(ValueError):
            self.base.retrieve("A", "lecture", "database")

    def test_course_and_material_filter_before_embedding_or_reranking(self):
        self.write("B/lecture/private.md", "Database secret for another course")
        self.write("A/tutorial/tute.md", "Database tutorial only")
        self.base.sync(self.config.knowledge_base_dir)
        rows = self.base.retrieve("A", "lecture", "database")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["path"], "A/lecture/Week 3.md")
        with self.assertRaises(ValueError):
            self.base.retrieve("unknown", "lecture", "database")

    def test_valid_citation_has_location_quote_and_separate_evidence(self):
        self.base.sync(self.config.knowledge_base_dir)
        answer = self.base.chat("A", "lecture", "Explain database transactions")
        self.assertIn("[S1]", answer.answer)
        self.assertEqual(answer.citations[0].week, 3)
        self.assertEqual(answer.citations[0].page, 1)
        self.assertTrue(answer.citations[0].quote)
        self.assertEqual(answer.evidence[0].rerank_score, 3)
        self.assertGreater(answer.estimated_cost_usd, 0)
        log = self.config.telemetry_path.read_text()
        self.assertNotIn("Database transactions", log)
        self.assertNotIn("Week 3", log)

    def test_invented_quote_and_source_id_fail_closed(self):
        self.base.sync(self.config.knowledge_base_dir)
        for flag in ("invalid_quote", "invalid_id"):
            setattr(self.client, flag, True)
            response = self.base.chat("A", "lecture", "Explain database transactions")
            self.assertEqual(response.citations, [])
            self.assertIn("引用校验", response.answer)
            self.assertEqual(response.diagnostics.reason,
                             "quote_not_found" if flag == "invalid_quote" else "unknown_source_id")
            setattr(self.client, flag, False)

    def test_low_relevance_abstains_without_generation(self):
        self.base.sync(self.config.knowledge_base_dir)
        self.client.low_relevance = True
        response = self.base.chat("A", "lecture", "What is the database exam room?")
        self.assertEqual(response.evidence, [])
        self.assertEqual(self.client.model_calls, 1)
        self.assertIn("没有足够证据", response.answer)

    def test_rejects_reranker_ids_not_supplied_by_server(self):
        self.base.sync(self.config.knowledge_base_dir)
        self.client.invalid_rerank = True
        with self.assertRaises(ValueError):
            self.base.retrieve("A", "all", "database")

    def test_embedding_configuration_change_requires_reindex(self):
        self.base.sync(self.config.knowledge_base_dir)
        config = self.config.model_copy(update={"rag_embedding_model": "another-model"})
        other = HybridKnowledgeBase(config, self.client)
        self.assertEqual(other.manifest(), {})
        self.assertEqual(other.sync(config.knowledge_base_dir).uploaded, 1)

    def test_source_api_and_chat_persist_page_citations(self):
        self.base.sync(self.config.knowledge_base_dir)
        app.dependency_overrides[get_settings] = lambda: self.config
        app.dependency_overrides[get_hybrid_base] = lambda: self.base
        app.dependency_overrides[get_knowledge_base] = lambda: SimpleNamespace()
        client = TestClient(app)
        answer = client.post("/api/chat", json={"course": "A", "material_type": "lecture",
                                               "question": "Explain database", "engine": "hybrid"})
        self.assertEqual(answer.status_code, 200, answer.text)
        payload = answer.json()
        chunk_id = payload["citations"][0]["chunk_id"]
        source = client.get(f"/api/sources/{chunk_id}")
        self.assertEqual(source.status_code, 200)
        self.assertEqual(source.json()["page"], 1)
        self.assertNotIn("sha256", source.json())
        messages = client.get(f"/api/conversations/{payload['conversation_id']}/messages").json()
        self.assertEqual(messages[-1]["citations"][0]["chunk_id"], chunk_id)
        file_response = client.get(f"/api/sources/{chunk_id}/file")
        self.assertEqual(file_response.status_code, 200)
        self.assertTrue(file_response.headers["content-disposition"].startswith("attachment"))
        self.path.write_text("Changed after indexing")
        self.assertEqual(client.get(f"/api/sources/{chunk_id}").status_code, 404)
        self.assertEqual(client.get(f"/api/sources/{chunk_id}/file").status_code, 404)

    def test_hybrid_api_rejects_followup_with_different_course(self):
        self.base.sync(self.config.knowledge_base_dir)
        app.dependency_overrides[get_settings] = lambda: self.config
        app.dependency_overrides[get_hybrid_base] = lambda: self.base
        app.dependency_overrides[get_knowledge_base] = lambda: SimpleNamespace()
        client = TestClient(app)
        answer = client.post("/api/chat", json={"course": "A", "material_type": "lecture",
                                               "question": "Explain database", "engine": "hybrid"})
        calls = self.client.model_calls
        followup = client.post("/api/chat", json={
            "course": "B", "material_type": "lecture", "question": "Tell me more",
            "conversation_id": answer.json()["conversation_id"], "engine": "hybrid",
        })
        self.assertEqual(followup.status_code, 409)
        self.assertEqual(calls, self.client.model_calls)

    def test_external_symlink_is_never_discovered(self):
        outside = Path(self.temp.name) / "outside.txt"
        outside.write_text("private")
        (self.path.parent / "escape.txt").symlink_to(outside)
        self.assertEqual(len(discover_documents(self.config.knowledge_base_dir)), 1)


class RetrievalTests(unittest.TestCase):
    def test_bm25_finds_exact_terms_and_chinese_bigrams(self):
        scores = bm25("ACID", ["ACID transactions", "nearest neighbors"])
        self.assertGreater(scores[0], scores[1])
        self.assertGreater(bm25("数据库", ["数据库事务", "机器学习"])[0], 0)

    def test_fusion_keeps_vector_only_matches_and_exact_keyword_matches(self):
        rows = [{"chunk_id": "a", "text": "ACID", "embedding": [0., 1.]},
                {"chunk_id": "b", "text": "transaction consistency", "embedding": [1., 0.]}]
        ranked = rank_candidates("ACID", rows, [1., 0.], 2)
        self.assertEqual({row["chunk_id"] for row in ranked}, {"a", "b"})
        self.assertGreater(next(r for r in ranked if r["chunk_id"] == "a")["bm25_score"], 0)
        self.assertEqual(next(r for r in ranked if r["chunk_id"] == "b")["vector_score"], 1)

    def test_cosine_rejects_dimension_mismatch(self):
        with self.assertRaises(ValueError):
            cosine([1], [1, 2])


class ParserTests(unittest.TestCase):
    def test_pdf_page_numbers_and_text(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.pdf"
            writer = PdfWriter()
            for word in ("First", "Second"):
                page = writer.add_blank_page(width=400, height=400)
                font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                         NameObject("/Subtype"): NameObject("/Type1"),
                                         NameObject("/BaseFont"): NameObject("/Helvetica")})
                page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
                    DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
                stream = DecodedStreamObject()
                stream.set_data(f"BT /F1 12 Tf 20 200 Td ({word} page) Tj ET".encode())
                page[NameObject("/Contents")] = writer._add_object(stream)
            writer.write(path)
            units = extract_units(path)
            self.assertEqual([unit.page for unit in units], [1, 2])
            self.assertIn("First", units[0].text)
            self.assertIn("Second", units[1].text)

    def test_pptx_preserves_slide_order_and_tables(self):
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.pptx"
            presentation = Presentation()
            for word in ("First", "Second"):
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = word
                table = slide.shapes.add_table(1, 1, 0, Inches(2), Inches(2), Inches(1)).table
                table.cell(0, 0).text = "Table evidence"
            presentation.save(path)
            units = extract_units(path)
            self.assertEqual([unit.locator_type for unit in units], ["slide", "slide"])
            self.assertEqual([unit.page for unit in units], [1, 2])
            self.assertIn("Second", units[1].text)
            self.assertIn("Table evidence", units[1].text)

    def test_image_only_page_is_explicit_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "A/lecture/blank.pdf"
            path.parent.mkdir(parents=True)
            path.write_bytes(b"%PDF-test")
            document = discover_documents(root)[0]
            with patch("app.services.parsing.extract_units",
                       return_value=[TextUnit(1, "page", "")]):
                with self.assertRaisesRegex(ValueError, "need OCR"):
                    chunk_document(document)

    def test_docx_uses_section_instead_of_fake_page_number(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.docx"
            document = Document()
            document.add_paragraph("Course notes")
            document.save(path)
            self.assertEqual(extract_units(path)[0].locator_type, "section")


if __name__ == "__main__":
    unittest.main()
